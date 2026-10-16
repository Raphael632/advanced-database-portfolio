"""Fetch verified NBS sources, build the database, and validate source-to-database fidelity.

Run from any directory. Source data stays outside Git in --workspace (default .local).
No workbook is edited, no source price is invented, and no archive is extracted to disk.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sqlite3
import sys
import tempfile
from urllib.request import Request, urlopen
from zipfile import ZipFile

import openpyxl

PREFIX = 'Ezenduka_Chukwubuikem_301913360'
HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
ZONES = ['NORTH CENTRAL', 'NORTH EAST', 'NORTH WEST', 'SOUTH EAST', 'SOUTH SOUTH', 'SOUTH WEST']
MONTHS = {name: n for n, names in enumerate([
    ['jan', 'january'], ['feb', 'february'], ['mar', 'march'], ['apr', 'april'],
    ['may'], ['jun', 'june'], ['jul', 'july'], ['aug', 'august'],
    ['sep', 'sept', 'september'], ['oct', 'october'], ['nov', 'november'], ['dec', 'december']
], 1) for name in names}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def manifest():
    return json.loads((HERE / 'sources.json').read_text(encoding='utf-8'))


def clean_label(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'Invalid item label: {value!r}')
    return re.sub(r'\s+', ' ', value).strip().rstrip(',').strip()


def reporting_period(header):
    match = re.fullmatch(r'Average of\s+([A-Za-z]+)\.?-(\d{2}|\d{4})', str(header).strip(), re.I)
    if not match or match[1].lower() not in MONTHS:
        raise ValueError(f'Unrecognised current-period header: {header!r}')
    year = int(match[2]) + (2000 if len(match[2]) == 2 else 0)
    return f'{year:04d}-{MONTHS[match[1].lower()]:02d}'


def kobo(value):
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError(f'Missing or non-numeric price: {value!r}')
    number = Decimal(str(value))
    if not number.is_finite() or number <= 0:
        raise ValueError(f'Invalid price: {value!r}')
    result = int((number * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    if result <= 0 or result > 2**63 - 1:
        raise ValueError('Price outside the positive SQLite integer range')
    return result


def fetch(workspace):
    raw = workspace / 'raw'
    raw.mkdir(parents=True, exist_ok=True)
    for source in manifest()['resources']:
        path = raw / f"resource_{source['resource_id']}.bin"
        if path.exists():
            data = path.read_bytes()
        else:
            request = Request(source['url'], headers={'User-Agent': 'MIT8103-Academic-Portfolio/1.0'})
            with urlopen(request, timeout=90) as response:
                data = response.read()
        if sha(data) != source['sha256']:
            raise ValueError(f"Source checksum changed for {source['resource_id']}; review before updating the manifest")
        if not path.exists():
            path.write_bytes(data)
        print(f"Verified resource {source['resource_id']} ({len(data)} bytes)")


def read_publications(workspace):
    """Validate all sources before building. Return local-only prices and provenance."""
    publications, resource_rows, warnings = [], [], []
    for source in manifest()['resources']:
        path = workspace / 'raw' / f"resource_{source['resource_id']}.bin"
        data = path.read_bytes()
        if sha(data) != source['sha256']:
            raise ValueError(f'Checksum mismatch: {path.name}')
        resource_rows.append((source['resource_id'], 1, source['url'], sha(data),
                              datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(), len(data)))
        with ZipFile(BytesIO(data)) as archive:
            if '[Content_Types].xml' in archive.namelist():
                files = [(f"NBS_resource_{source['resource_id']}.xlsx", data)]
            else:
                files = [(n, archive.read(n)) for n in archive.namelist() if n.lower().endswith('.xlsx')]
        found = []
        for filename, content in files:
            workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
            try:
                national_sheets = [s for s in workbook if str(s.cell(1, 4).value).startswith('Average of')]
                if len(national_sheets) != 1 or 'Zone All item' not in workbook.sheetnames:
                    raise ValueError(f'Unexpected workbook structure: {filename}')
                national = national_sheets[0]
                period = reporting_period(national.cell(1, 4).value)
                if period not in source['periods']:
                    raise ValueError(f'Unexpected period {period} in resource {source["resource_id"]}')
                found.append(period)
                if period in ('2025-03', '2025-04'):
                    warnings.append({'period': period, 'issue': 'Stale worksheet title; period taken from current-month column header', 'source_sheet': national.title})
                zone = workbook['Zone All item']
                if [zone.cell(1, c).value for c in range(2, 8)] != ZONES:
                    raise ValueError(f'Unexpected regional header order in {filename}')
                observations = []
                for sheet, columns in [(national, [(4, 'Nigeria')]), (zone, list(enumerate(ZONES, 2)))]:
                    seen = set()
                    for rownum, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if rownum == 1:
                            continue
                        label = row[0]
                        if label is None:
                            if any(row[c-1] is not None for c, _ in columns):
                                raise ValueError(f'Price without item label at {sheet.title}!{rownum}')
                            continue
                        cleaned = clean_label(label)
                        if cleaned in seen:
                            raise ValueError(f'Duplicate cleaned label: {cleaned}')
                        seen.add(cleaned)
                        for col, geography in columns:
                            observations.append({'item': cleaned, 'geography': geography, 'price_kobo': kobo(row[col-1]),
                                                 'source_sheet': sheet.title,
                                                 'source_cell': f'{openpyxl.utils.get_column_letter(col)}{rownum}'})
                    if len(seen) != 42:
                        raise ValueError(f'Expected 42 source items in {sheet.title}; got {len(seen)}')
                publications.append({'resource_id': source['resource_id'], 'period': period,
                                     'workbook_name': filename, 'workbook_sha256': sha(content),
                                     'observations': observations})
            finally:
                workbook.close()
        if sorted(found) != sorted(source['periods']):
            raise ValueError(f'Missing/duplicate workbooks for resource {source["resource_id"]}')
    if sorted(p['period'] for p in publications) != [f'2025-{m:02d}' for m in range(1, 13)]:
        raise ValueError('Expected each month of 2025 exactly once')
    warnings.append({'issue': 'National Smoked fish (Mackerel) and zonal Smoked fish kept as separate item labels; equivalence not established'})
    return sorted(publications, key=lambda x: x['period']), resource_rows, warnings


def connect(path):
    connection = sqlite3.connect(path)
    connection.execute('PRAGMA foreign_keys = ON')
    return connection


def build(workspace):
    database = workspace / f'{PREFIX}_Food_Prices.sqlite3'
    if database.exists():
        raise FileExistsError(f'Database exists: {database}. Use validate, or build in a new workspace to preserve review work.')
    publications, resources, warnings = read_publications(workspace)
    workspace.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=workspace, suffix='.sqlite3', delete=False) as temporary:
        pending = Path(temporary.name)
    connection = connect(pending)
    try:
        connection.executescript((HERE / 'schema.sql').read_text(encoding='utf-8'))
        with connection:
            connection.execute('INSERT INTO publisher VALUES (1,?,?)', ('National Bureau of Statistics, Nigeria', 'https://www.nigerianstat.gov.ng/'))
            connection.executemany('INSERT INTO source_resource VALUES (?,?,?,?,?,?)', resources)
            labels = sorted({o['item'] for p in publications for o in p['observations']})
            items = {label: n for n, label in enumerate(labels, 1)}
            geographies = {label: n for n, label in enumerate(['Nigeria'] + ZONES, 1)}
            connection.executemany('INSERT INTO item VALUES (?,?)', [(i, label) for label, i in items.items()])
            connection.executemany('INSERT INTO geography VALUES (?,?,?)', [(i, label, 'national' if label == 'Nigeria' else 'zone') for label, i in geographies.items()])
            for n, publication in enumerate(publications, 1):
                connection.execute('INSERT INTO publication VALUES (?,?,?,?,?)',
                                   (n, publication['resource_id'], publication['period'], publication['workbook_name'], publication['workbook_sha256']))
                connection.executemany('INSERT INTO price_observation(publication_id,item_id,geography_id,price_kobo,source_sheet,source_cell) VALUES (?,?,?,?,?,?)',
                                       [(n, items[o['item']], geographies[o['geography']], o['price_kobo'], o['source_sheet'], o['source_cell']) for o in publication['observations']])
            connection.executemany('INSERT INTO app_user VALUES (?,?,?,1)', [(1, 'demo_analyst_a', 'analyst'), (2, 'demo_steward_b', 'steward')])
        if connection.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
            raise ValueError('SQLite integrity check failed')
        connection.close()
        pending.replace(database)
    except BaseException:
        connection.close()
        pending.unlink(missing_ok=True)
        raise
    (workspace / 'source_warnings.json').write_text(json.dumps(warnings, indent=2), encoding='utf-8')
    print(f'Built {database}')
    return validate(workspace)


def validate(workspace):
    publications, resources, warnings = read_publications(workspace)
    database = workspace / f'{PREFIX}_Food_Prices.sqlite3'
    if not database.is_file():
        raise FileNotFoundError(database)
    connection = connect(database)
    try:
        expected = {(p['period'], o['item'], o['geography']):
                    (o['price_kobo'], p['resource_id'], p['workbook_name'], o['source_sheet'], o['source_cell'])
                    for p in publications for o in p['observations']}
        rows = connection.execute('SELECT period,item,geography,price_kobo,resource_id,workbook_name,source_sheet,source_cell FROM published_prices').fetchall()
        actual = {tuple(r[:3]): tuple(r[3:]) for r in rows}
        if len(rows) != len(actual) or actual != expected:
            raise AssertionError('Source-to-database values or provenance do not reconcile')
        integrity = connection.execute('PRAGMA integrity_check').fetchone()[0]
        fk = connection.execute('PRAGMA foreign_key_check').fetchall()
        if integrity != 'ok' or fk:
            raise AssertionError('Database integrity failed')
        counts = {name: connection.execute(f'SELECT count(*) FROM {name}').fetchone()[0]
                  for name in ('publisher','source_resource','publication','item','geography','price_observation','app_user','review_task')}
        monthly = [{'period': p, 'observations': n} for p, n in connection.execute(
            'SELECT period,count(*) FROM published_prices GROUP BY period ORDER BY period')]
        report = {'checked_at_utc': datetime.now(timezone.utc).isoformat(),
                  'python_version': sys.version.split()[0], 'sqlite_version': sqlite3.sqlite_version,
                  'openpyxl_version': openpyxl.__version__, 'counts': counts, 'monthly': monthly,
                  'source_cells_reconciled': len(rows), 'integrity_check': integrity,
                  'foreign_key_violations': len(fk), 'price_rounding': 'NGN to integer kobo using decimal ROUND_HALF_UP',
                  'synthetic_prices': 0, 'workflow_users': '2 explicitly fictional demo accounts',
                  'warnings': warnings}
        (workspace / 'validation.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
        print(json.dumps(report, indent=2))
        return report
    finally:
        connection.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['fetch','build','validate','all'])
    parser.add_argument('--workspace', type=Path, default=ROOT / '.local')
    args = parser.parse_args()
    workspace = args.workspace.resolve()
    if args.command in ('fetch','all'):
        fetch(workspace)
    if args.command in ('build','all'):
        build(workspace)
    if args.command == 'validate':
        validate(workspace)


if __name__ == '__main__':
    main()
